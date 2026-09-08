"""Shark Tank–style prototype showcase deck (16:9)."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x1C, 0x19, 0x17)
MUTED = RGBColor(0x57, 0x53, 0x4E)
PAPER = RGBColor(0xF3, 0xEF, 0xE8)
SURFACE = RGBColor(0xFF, 0xFD, 0xF8)
ACCENT = RGBColor(0x5B, 0x21, 0xB6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x14, 0x11, 0x10)
CREAM = RGBColor(0xE7, 0xE5, 0xE4)
OK = RGBColor(0x16, 0x65, 0x34)
WARN = RGBColor(0x85, 0x4D, 0x0E)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, text, size=20, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_box(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(shape, color)
    return shape


def add_round(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(shape, color)
    shape.adjustments[0] = 0.08
    return shape


def add_text(slide, l, t, w, h, text, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return box


def bullets(slide, l, t, w, h, items, size=18, color=INK):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        set_run(p.add_run(), item, size, False, color)
    return box


def footer(slide, dark=False):
    if dark:
        return
    add_box(slide, 0, H - Inches(0.38), W, Inches(0.38), DARK)
    add_text(
        slide,
        Inches(0.5),
        H - Inches(0.34),
        Inches(12.3),
        Inches(0.28),
        "MERIDIAN TEST CLIENT   ·   Prototype showcase   ·   Class of 2025   ·   July 2026",
        11,
        False,
        WHITE,
    )


def light(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, PAPER)
    add_box(s, 0, 0, Inches(0.14), H, ACCENT)
    footer(s)
    return s


def dark(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, DARK)
    add_box(s, 0, 0, Inches(0.14), H, ACCENT)
    return s


def kicker(slide, text, y=0.32, color=ACCENT):
    add_text(slide, Inches(0.55), Inches(y), Inches(12.2), Inches(0.32), text.upper(), 13, True, color)


def headline(slide, text, y=0.7, size=36, color=INK, height=1.4):
    add_text(slide, Inches(0.55), Inches(y), Inches(12.2), Inches(height), text, size, True, color)


def card(slide, x, y, w, h, title, body, title_size=16, body_size=14):
    add_round(slide, Inches(x), Inches(y), Inches(w), Inches(h), SURFACE)
    add_text(slide, Inches(x + 0.28), Inches(y + 0.22), Inches(w - 0.5), Inches(0.45), title, title_size, True, ACCENT)
    add_text(slide, Inches(x + 0.28), Inches(y + 0.7), Inches(w - 0.5), Inches(h - 0.95), body, body_size, False, INK)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 HOOK
    s = dark(prs)
    add_text(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.4), "PROTOTYPE SHOWCASE", 14, True, ACCENT)
    add_text(
        s,
        Inches(0.6),
        Inches(2.05),
        Inches(12.1),
        Inches(2.4),
        "School employment reports\nstill start with two SAS programs\nand a shared output folder.",
        40,
        True,
        WHITE,
    )
    add_text(
        s,
        Inches(0.6),
        Inches(5.5),
        Inches(12.1),
        Inches(0.9),
        "We replaced the desktop ritual with a signed-in product — without inventing a single business rule.",
        20,
        False,
        CREAM,
    )

    # 2 NAME
    s = dark(prs)
    add_text(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(0.4), "WE BUILT", 14, True, ACCENT)
    add_text(s, Inches(0.6), Inches(2.3), Inches(12.1), Inches(1.2), "Meridian Test Client", 48, True, WHITE)
    add_text(
        s,
        Inches(0.6),
        Inches(3.7),
        Inches(12.1),
        Inches(1.6),
        "The reporting shop’s command center: import the graduate workbook, generate the same seven-page school PDF, download only what you are allowed to see, and ask questions that stay inside authorized evidence.",
        22,
        False,
        CREAM,
    )
    add_text(s, Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.4), ".NET 8   ·   SQLite   ·   React   ·   QuestPDF   ·   Identity   ·   Knowledge Assistant", 16, False, RGBColor(0xA8, 0xA2, 0x9E))

    # 3 CUSTOMER
    s = light(prs)
    kicker(s, "The customer")
    headline(s, "A reporting shop that must ship one PDF per school, every cycle.")
    card(s, 0.55, 2.4, 3.95, 4.0, "They receive", "An Excel extract of graduates: employment, demographics, employer type, region, salary.")
    card(s, 4.7, 2.4, 3.95, 4.0, "They must produce", "One Class of 2025 letter PDF per school. Seven pages. Same tables the schools already recognize.")
    card(s, 8.85, 2.4, 3.95, 4.0, "They cannot break", "Salary suppression. Recodes. Totals. Who is allowed to open which file.")

    # 4 PAIN
    s = light(prs)
    kicker(s, "The old way")
    headline(s, "Yesterday’s “product” was a specialist and a folder.")
    pains = [
        ("Install SAS 9.4", "Point two programs at local paths. Hope the folders exist."),
        ("Run the black box", "createschrptfiles2025.sas, then schreptsummary_2025.sas."),
        ("Collect PDFs", "ODS print layout. Shared output. No sign-in."),
        ("Email the files", "Anyone with the folder sees every school."),
    ]
    for i, (t, b) in enumerate(pains):
        card(s, 0.55 + i * 3.15, 2.45, 3.0, 3.9, t, b, 16, 15)

    # 5 COST
    s = dark(prs)
    add_text(s, Inches(0.6), Inches(0.55), Inches(12.1), Inches(0.35), "WHAT THAT COST", 14, True, ACCENT)
    add_text(s, Inches(0.6), Inches(1.1), Inches(12.1), Inches(1.1), "If only one person can run SAS,\nthe shop cannot scale — or prove — a cycle.", 32, True, WHITE)
    lines = [
        "No dashboard of what is actually stored.",
        "No run history when a school asks “did you generate mine?”",
        "No test that salary stats stay hidden when n < 5.",
        "No 404 when the wrong person guesses a download link.",
        "A print PDF: only the first school’s ods pdf used the accessible option.",
    ]
    bullets(s, Inches(0.6), Inches(2.6), Inches(12.0), Inches(4.2), lines, 22, CREAM)

    # 6 PITCH
    s = light(prs)
    kicker(s, "The pitch")
    headline(s, "Same SAS rules. A product instead of a ritual.")
    add_text(
        s,
        Inches(0.55),
        Inches(2.3),
        Inches(12.2),
        Inches(1.1),
        "Meridian Test Client is a local web app that runs the characterized school-employment pipeline — import, calculate, generate, authorize, answer — without rewriting the business logic to look cleaner.",
        20,
    )
    card(s, 0.55, 3.7, 4.0, 2.7, "Keep", "CF-S-00. Seven pages. Gray headers. “.” for missing. Class of 2025 / July 2026.")
    card(s, 4.7, 3.7, 4.0, 2.7, "Add", "Sign-in, school grants, run history, year dashboard, tagged PDF, Adobe-checked reading order.")
    card(s, 8.85, 3.7, 4.0, 2.7, "Refuse", "Invented suppression. Restyled reports. A second database. The assistant as the calculator.")

    # 7 HOW IT WORKS
    s = light(prs)
    kicker(s, "How it works")
    headline(s, "Six clicks. One cycle.")
    steps = [
        ("1", "Sign in", "Identity cookie. Admin, report user, or viewer."),
        ("2", "Import", "Admin drops the .xlsx. Invalid rows logged. Duplicate hash rejected."),
        ("3", "See stock", "Dashboard: schools and graduates by class year from stored rows."),
        ("4", "Generate", "One school or all. Sequential or bounded parallel (1–8)."),
        ("5", "Deliver", "Authorized /downloads/reports/{id}. Else 404."),
        ("6", "Ask", "Assistant quotes printed PDFs and docs. It does not rerun SAS."),
    ]
    for i, (n, t, b) in enumerate(steps):
        x = 0.45 + (i % 3) * 4.25
        y = 2.25 + (i // 3) * 2.25
        add_round(s, Inches(x), Inches(y), Inches(4.05), Inches(2.05), SURFACE)
        add_text(s, Inches(x + 0.22), Inches(y + 0.18), Inches(0.5), Inches(0.4), n, 22, True, ACCENT)
        add_text(s, Inches(x + 0.7), Inches(y + 0.22), Inches(3.1), Inches(0.4), t, 18, True, INK)
        add_text(s, Inches(x + 0.22), Inches(y + 0.75), Inches(3.6), Inches(1.1), b, 14, False, MUTED)

    # 8 THE REPORT
    s = light(prs)
    kicker(s, "The artifact schools already know")
    headline(s, "Seven letter pages. We did not redesign the report.")
    card(s, 0.55, 2.4, 6.05, 4.0, "What they open", "Page 1: Total Reported under the headers. Employment, gender, race, employer type, region, salaries. Note, then test-client footer in the flow. Missing cells stay “.” — screen-reader name: Not displayed.")
    card(s, 6.8, 2.4, 6.0, 4.0, "What we added behind it", "Tags, language, reading order. Adobe Acrobat accessibility checker: passed, screen order correct. That check is recorded separately. Not a veraPDF / PAC / NVDA pack.")

    # 9 ACCESS
    s = light(prs)
    kicker(s, "Trust")
    headline(s, "The wrong person never “finds” another school.")
    rows = [
        ("Admin", "Loads the cycle. Can generate every school."),
        ("Report user", "Generates a school they are granted."),
        ("Viewer", "Reads granted reports. Cannot import. Cannot generate all."),
        ("No grant", "Guessing a report id returns 404 — no name, no path."),
    ]
    for i, (t, b) in enumerate(rows):
        y = 2.3 + i * 1.05
        add_round(s, Inches(0.55), Inches(y), Inches(12.2), Inches(0.95), SURFACE)
        add_text(s, Inches(0.85), Inches(y + 0.25), Inches(2.6), Inches(0.45), t, 18, True, ACCENT)
        add_text(s, Inches(3.6), Inches(y + 0.25), Inches(8.8), Inches(0.5), b, 18, False, INK)

    # 10 ASSISTANT
    s = light(prs)
    kicker(s, "The supporting act")
    headline(s, "Ask the reports. Do not ask the model to invent totals.")
    add_text(
        s,
        Inches(0.55),
        Inches(2.3),
        Inches(12.2),
        Inches(1.0),
        "The Knowledge Assistant searches authorized catalog text and generated PDFs. It quotes printed figures. It can sum printed Total Reported. It never recalculates SAS, salaries, or suppression from raw rows.",
        20,
    )
    card(s, 0.55, 3.6, 4.0, 2.8, "This report", "Locked to one PDF you can view.")
    card(s, 4.7, 3.6, 4.0, 2.8, "All reports you can see", "Sums, last year, difference — from printed numbers.")
    card(s, 8.85, 3.6, 4.0, 2.8, "No evidence", "We show Insufficient evidence. We do not guess.")

    # 11 SECRET
    s = dark(prs)
    add_text(s, Inches(0.6), Inches(0.55), Inches(12.1), Inches(0.35), "THE UNFAIR ADVANTAGE", 14, True, ACCENT)
    add_text(s, Inches(0.6), Inches(1.05), Inches(12.1), Inches(1.3), "We characterized SAS before we wrote C#.", 34, True, WHITE)
    add_text(
        s,
        Inches(0.6),
        Inches(2.5),
        Inches(12.1),
        Inches(1.1),
        "Every count, recode, and suppression has a Rule ID. Eighteen ambiguous cases stayed skipped. We did not “clean up” the rules to make a prettier demo.",
        20,
        False,
        CREAM,
    )
    add_round(s, Inches(0.6), Inches(4.0), Inches(12.1), Inches(2.4), RGBColor(0x2A, 0x22, 0x1B))
    add_text(s, Inches(0.95), Inches(4.25), Inches(11.4), Inches(0.45), "CF-S-00 — the line investors should remember", 18, True, ACCENT)
    add_text(
        s,
        Inches(0.95),
        Inches(4.8),
        Inches(11.4),
        Inches(1.3),
        "Salary statistics print only when n ≥ 5 on non-missing full-time long-term salaries. Not on headcount. An AI proposal to emit salaries when n < 5 was rejected. The test still fails that change on purpose.",
        18,
        False,
        CREAM,
    )

    # 12 PROOF
    s = light(prs)
    kicker(s, "Traction — what we can show today")
    headline(s, "A running prototype. Tests. A human veto log.")
    proofs = [
        ("Live app", "Sign-in, import, generate, history, assistant — https://localhost:7117"),
        ("Locked rules", "Characterization + calculator tests. CF-S-00 cannot silently change."),
        ("Adobe check", "Accessibility checker passed. Reading / screen order correct."),
        ("CI gate", "GitHub Actions builds and tests on main and every PR."),
        ("Authz proof", "Unauthorized download is 404. RAG filters before the model."),
        ("Honest miss", "Baseline PDF vs sample Excel is a known population mismatch. We did not chase it."),
    ]
    for i, (t, b) in enumerate(proofs):
        x = 0.55 + (i % 3) * 4.2
        y = 2.35 + (i // 3) * 2.15
        card(s, x, y, 4.0, 2.0, t, b, 16, 14)

    # 13 VS SAS
    s = light(prs)
    kicker(s, "Why not keep SAS?")
    headline(s, "SAS still knows the rules. It is not a product.")
    add_round(s, Inches(0.55), Inches(2.35), Inches(6.05), Inches(4.1), SURFACE)
    add_text(s, Inches(0.85), Inches(2.55), Inches(5.5), Inches(0.4), "SAS shop", 18, True, MUTED)
    bullets(s, Inches(0.85), Inches(3.1), Inches(5.5), Inches(3.0), [
        "License + local paths",
        "Rules untested",
        "Print PDF",
        "Shared folder = access control",
        "Staff memory = knowledge base",
    ], 18, INK)
    add_round(s, Inches(6.8), Inches(2.35), Inches(6.0), Inches(4.1), RGBColor(0xF3, 0xE8, 0xFF))
    add_text(s, Inches(7.1), Inches(2.55), Inches(5.5), Inches(0.4), "Meridian Test Client", 18, True, ACCENT)
    bullets(s, Inches(7.1), Inches(3.1), Inches(5.5), Inches(3.0), [
        "Browser + SQLite, no SAS to run a cycle",
        "Rule IDs + automated tests",
        "Same look, tagged, Adobe-checked",
        "Roles and school grants on the server",
        "Assistant on authorized evidence only",
    ], 18, INK)

    # 14 DEMO
    s = dark(prs)
    add_text(s, Inches(0.6), Inches(0.7), Inches(12.1), Inches(0.35), "LIVE PROTOTYPE", 14, True, ACCENT)
    add_text(s, Inches(0.6), Inches(1.2), Inches(12.1), Inches(1.0), "Don’t take the slides.\nTake the product.", 40, True, WHITE)
    items = [
        "Admin signs in. Dashboard shows stored graduates by year.",
        "Import is Admin-only. A repeat file is rejected.",
        "Generate one school. Open the seven-page PDF.",
        "Ask: “What is Total Reported in this report?”",
        "Sign in as Viewer. Import denied. Wrong PDF id → 404.",
    ]
    bullets(s, Inches(0.6), Inches(3.4), Inches(12.0), Inches(3.4), items, 20, CREAM)

    # 15 HONEST
    s = light(prs)
    kicker(s, "The fine print investors should hear")
    headline(s, "We will not oversell the prototype.")
    card(s, 0.55, 2.4, 6.05, 4.0, "What is true", "Local MVP. SQLite only. Totals from C#. Adobe Acrobat checker observed on generated PDFs. Human rejected “print salaries when n < 5.”")
    card(s, 6.8, 2.4, 6.0, 4.0, "What is not claimed", "Cloud scale. A second database. veraPDF / PAC / NVDA pack. Playwright CI. Matching the Test University baseline headcount. The assistant as source of truth.")

    # 16 ASK
    s = light(prs)
    kicker(s, "The ask")
    headline(s, "Green-light the prototype. Hold the rules.")
    add_text(
        s,
        Inches(0.55),
        Inches(2.3),
        Inches(12.2),
        Inches(1.0),
        "This capstone is ready to demo as a shop tool. The next money — if it ever comes — is validation packs, a real identity provider, and performance on a full generate-all. Not a rewrite of CF-S-00.",
        20,
    )
    card(s, 0.55, 3.6, 4.0, 2.8, "Use it", "Run a cycle on the sample workbook. Generate one school. Download as Admin and fail as Viewer.")
    card(s, 4.7, 3.6, 4.0, 2.8, "Judge it", "Did we preserve SAS? Did we lock suppression? Did we stop the wrong download?")
    card(s, 8.85, 3.6, 4.0, 2.8, "Don’t ask us to", "Restyle the PDF. Guess skipped SAS. Put Excel rows into the assistant index.")

    # 17 CLOSE
    s = dark(prs)
    add_text(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(0.4), "MERIDIAN TEST CLIENT", 14, True, ACCENT)
    add_text(s, Inches(0.6), Inches(2.35), Inches(12.1), Inches(2.0), "The reports did not change.\nThe shop finally can.", 40, True, WHITE)
    add_text(
        s,
        Inches(0.6),
        Inches(5.0),
        Inches(12.1),
        Inches(1.2),
        "Questions.\nClass of 2025  ·  July 2026  ·  docs/capstone-demo/",
        18,
        False,
        CREAM,
    )

    folder = r"c:\Users\104046\Desktop\AccessibleSchoolReportModernizer\docs\capstone-demo"
    out = folder + r"\Accessible-School-Report-Modernizer-Capstone-Demo.pptx"
    fallback = folder + r"\Accessible-School-Report-Modernizer-Capstone-Demo-SHOWCASE.pptx"
    try:
        prs.save(out)
        print(out, "slides", len(prs.slides))
    except PermissionError:
        prs.save(fallback)
        print(fallback, "slides", len(prs.slides), "(original pptx is open)")


if __name__ == "__main__":
    build()
